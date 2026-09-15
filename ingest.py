"""
QPFP Genius — course corpus ingestion.

Runs LOCALLY on the machine that holds the Network FP course folder. Walks every
Handout, PPT and ProTool workbook, extracts the text, tags it by level/session/kind,
chunks it, and writes data/course_corpus.json — which the MCP server then serves.

The server itself never needs these libraries; only this script does.

Usage:
    pip install -r requirements-ingest.txt
    python3 ingest.py "/Users/praxysmac/Downloads/QPFP B15"

Then:
    git add data/course_corpus.json && git commit -m "Course corpus" && git push
"""

from __future__ import annotations

import json
import re
import sys
from datetime import datetime, timezone
from pathlib import Path

CHUNK_CHARS = 1400
CHUNK_OVERLAP = 180
MIN_CHUNK = 60

SKIP_DIRS = {"_source", "__MACOSX", ".git", "node_modules"}
SKIP_FILE_PREFIXES = ("~$", ".")


# ------------------------------------------------------------------ extract

def extract_pdf(p: Path) -> str:
    from pypdf import PdfReader
    out = []
    reader = PdfReader(str(p))
    for i, page in enumerate(reader.pages, 1):
        try:
            t = page.extract_text() or ""
        except Exception:
            t = ""
        if t.strip():
            out.append(f"[page {i}]\n{t}")
    return "\n\n".join(out)


def extract_docx(p: Path) -> str:
    import docx
    d = docx.Document(str(p))
    parts = [para.text for para in d.paragraphs if para.text.strip()]
    for tbl in d.tables:
        for row in tbl.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def extract_pptx(p: Path) -> str:
    from pptx import Presentation
    prs = Presentation(str(p))
    out = []
    for i, slide in enumerate(prs.slides, 1):
        bits = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                bits.append(shape.text_frame.text.strip())
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    cells = [c.text.strip() for c in row.cells]
                    if any(cells):
                        bits.append(" | ".join(cells))
        try:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                bits.append(f"[speaker notes] {notes}")
        except Exception:
            pass
        if bits:
            out.append(f"[slide {i}]\n" + "\n".join(bits))
    return "\n\n".join(out)


def extract_xlsx(p: Path) -> str:
    import openpyxl
    out = []
    for data_only in (False, True):
        try:
            wb = openpyxl.load_workbook(str(p), data_only=data_only, read_only=True)
        except Exception:
            continue
        tag = "values" if data_only else "formulas"
        for ws in wb.worksheets:
            rows = []
            for r in ws.iter_rows():
                cells = []
                for c in r:
                    if c.value not in (None, ""):
                        cells.append(f"{c.coordinate}={c.value}")
                if cells:
                    rows.append(" · ".join(cells))
                if len(rows) > 400:
                    break
            if rows:
                out.append(f"[sheet {ws.title} · {tag}]\n" + "\n".join(rows))
        try:
            wb.close()
        except Exception:
            pass
        break  # formulas pass is enough; values pass only if formulas failed
    return "\n\n".join(out)


EXTRACTORS = {
    ".pdf": extract_pdf,
    ".docx": extract_docx,
    ".doc": None,
    ".pptx": extract_pptx,
    ".ppt": None,
    ".xlsx": extract_xlsx,
    ".xlsm": extract_xlsx,
    ".md": lambda p: p.read_text(encoding="utf-8", errors="ignore"),
    ".txt": lambda p: p.read_text(encoding="utf-8", errors="ignore"),
}


# --------------------------------------------------------------- classify

LEVEL_RE = re.compile(r"Level\s+(I{1,3})\b", re.I)
SESSION_RE = re.compile(r"(?:session|handout|ppt|protool)\D{0,4}(\d{1,2})\b", re.I)
LEADING_NUM_RE = re.compile(r"\b(\d{1,2})\s*[-–—.]\s*")

LEVEL_NAME = {"I": "Level I", "II": "Level II", "III": "Level III"}


def classify(path: Path, root: Path) -> tuple[str | None, int | None, str]:
    rel = str(path.relative_to(root))
    low = rel.lower()

    level = None
    m = LEVEL_RE.search(rel)
    if m:
        level = LEVEL_NAME.get(m.group(1).upper())

    session = None
    for candidate in (path.name, rel):
        m = SESSION_RE.search(candidate)
        if m:
            n = int(m.group(1))
            if 1 <= n <= 36:
                session = n
                break
    if session is None:
        m = LEADING_NUM_RE.search(path.parent.name)
        if m and 1 <= int(m.group(1)) <= 36:
            session = int(m.group(1))

    if "handout" in low:
        kind = "handout"
    elif "ppt" in low or path.suffix.lower() in {".pptx", ".ppt"}:
        kind = "ppt"
    elif "protool" in low or path.suffix.lower() in {".xlsx", ".xlsm"}:
        kind = "protool"
    elif "assessment" in low or "quiz" in low or "exam" in low:
        kind = "assessment"
    else:
        kind = "other"

    if level is None and session is not None:
        level = "Level I" if session <= 12 else "Level II" if session <= 24 else "Level III"

    return level, session, kind


# ----------------------------------------------------------------- chunk

def chunk(text: str) -> list[str]:
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if not text:
        return []
    paras = [p.strip() for p in text.split("\n\n") if p.strip()]
    chunks, buf = [], ""
    for p in paras:
        if len(buf) + len(p) + 2 <= CHUNK_CHARS:
            buf = f"{buf}\n\n{p}" if buf else p
            continue
        if buf:
            chunks.append(buf)
        while len(p) > CHUNK_CHARS:
            cut = p.rfind("\n", 0, CHUNK_CHARS)
            cut = cut if cut > CHUNK_CHARS // 2 else CHUNK_CHARS
            chunks.append(p[:cut])
            p = p[max(0, cut - CHUNK_OVERLAP):]
        buf = p
    if buf:
        chunks.append(buf)
    return [c for c in chunks if len(c) >= MIN_CHUNK]


# ------------------------------------------------------------------ main

def main() -> int:
    if len(sys.argv) < 2:
        print(__doc__)
        return 1
    root = Path(sys.argv[1]).expanduser()
    if not root.is_dir():
        print(f"Not a directory: {root}")
        return 1

    out_path = Path(__file__).parent / "data" / "course_corpus.json"
    docs, skipped, failed = [], [], []

    files = [
        p for p in sorted(root.rglob("*"))
        if p.is_file()
        and not p.name.startswith(SKIP_FILE_PREFIXES)
        and not any(part in SKIP_DIRS for part in p.parts)
    ]

    for p in files:
        ext = p.suffix.lower()
        fn = EXTRACTORS.get(ext)
        if fn is None:
            skipped.append(f"{p.relative_to(root)} ({ext or 'no ext'})")
            continue
        try:
            text = fn(p)
        except Exception as e:
            failed.append(f"{p.relative_to(root)} — {type(e).__name__}: {e}")
            continue
        chunks = chunk(text)
        if not chunks:
            failed.append(f"{p.relative_to(root)} — no extractable text")
            continue
        level, session, kind = classify(p, root)
        docs.append({
            "id": f"d{len(docs):04d}",
            "file": str(p.relative_to(root)),
            "title": p.stem,
            "level": level,
            "session": session,
            "kind": kind,
            "chunks": chunks,
        })
        print(f"  ✓ [{kind:10s} s{session if session else '--':>2}] "
              f"{p.relative_to(root)}  ({len(chunks)} chunks)")

    payload = {
        "generated": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "source_root": str(root),
        "doc_count": len(docs),
        "chunk_count": sum(len(d["chunks"]) for d in docs),
        "docs": docs,
    }
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")

    mb = out_path.stat().st_size / 1_048_576
    print("\n" + "=" * 62)
    print(f"Documents ingested : {len(docs)}")
    print(f"Chunks             : {payload['chunk_count']:,}")
    print(f"Output             : {out_path}  ({mb:.1f} MB)")
    sessions = sorted({d["session"] for d in docs if d["session"]})
    missing = [n for n in range(1, 37) if n not in sessions]
    print(f"Sessions covered   : {len(sessions)}/36")
    if missing:
        print(f"  ! No content for sessions: {missing}")
    if skipped:
        print(f"\nSkipped (unsupported format): {len(skipped)}")
        for s in skipped[:12]:
            print(f"  - {s}")
        print("  Convert legacy .doc/.ppt to .docx/.pptx and re-run to include them.")
    if failed:
        print(f"\nFailed to extract: {len(failed)}")
        for f in failed[:12]:
            print(f"  - {f}")
    if mb > 80:
        print("\n! Corpus exceeds 80 MB. Consider excluding ProTool workbooks "
              "(the Formula Bible already covers them) to keep the repo lean.")
    print("=" * 62)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
